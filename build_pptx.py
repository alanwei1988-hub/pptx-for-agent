"""
HTML → PNG → PPTX pipeline (v1.image strategy: full-bleed image embedding).

Usage:
    python build_pptx.py <input.html> [-o output.pptx]

Loads an HTML file containing one or more `.slide` elements (each sized
1920x1080), renders each via headless Chromium at 2x device scale (3840x2160
PNG), then assembles a 16:9 PPTX with each PNG embedded as a full-bleed
picture on a blank slide.

Why this design:
    Marp / template engines cap CSS expressiveness, so the LLM's design
    ability gets clipped at the framework boundary. By having the LLM emit
    free-form HTML/CSS+SVG and treating the rendered raster as the slide's
    final form, we expose the full LLM design ceiling at the cost of
    in-PowerPoint editability. Edits happen by re-prompting, not by clicking
    text boxes — the AI-native workflow.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

SLIDE_W_IN = 13.333  # 16:9 widescreen (PowerPoint default)
SLIDE_H_IN = 7.5
VIEWPORT = {"width": 1920, "height": 1080}
SCALE = 2  # device pixel ratio → 3840x2160 PNG for retina sharpness


def render_slides(html_path: Path, png_dir: Path) -> list[Path]:
    """Open html_path, screenshot each .slide element, return list of PNGs."""
    png_dir.mkdir(parents=True, exist_ok=True)
    out: list[Path] = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        ctx = browser.new_context(viewport=VIEWPORT, device_scale_factor=SCALE)
        page = ctx.new_page()
        page.goto(html_path.as_uri(), wait_until="networkidle")
        # Wait for web fonts to actually load before taking screenshots
        page.evaluate("() => document.fonts.ready")
        page.wait_for_timeout(500)  # tiny safety margin for font paint

        slides = page.locator(".slide")
        n = slides.count()
        if n == 0:
            raise SystemExit(f"no .slide elements found in {html_path}")
        print(f"  found {n} .slide elements")
        for i in range(n):
            target = png_dir / f"slide_{i + 1:02d}.png"
            slides.nth(i).screenshot(path=str(target), omit_background=False)
            out.append(target)
            print(f"  rendered {target.name}  ({target.stat().st_size // 1024} KB)")
        browser.close()
    return out


def build_pptx(png_paths: list[Path], pptx_path: Path) -> Path:
    """Assemble a 16:9 PPTX with one full-bleed picture per slide."""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    blank_layout = prs.slide_layouts[6]  # fully blank
    for png in png_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(png), 0, 0,
            width=prs.slide_width,
            height=prs.slide_height,
        )
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))
    print(f"  saved {pptx_path.name}  ({pptx_path.stat().st_size // 1024} KB)")
    return pptx_path


def main() -> None:
    ap = argparse.ArgumentParser(description="HTML → PPTX (v1.image)")
    ap.add_argument("html", type=Path, help="input HTML with .slide elements")
    ap.add_argument(
        "-o", "--output", type=Path, default=None,
        help="output .pptx path (default: build/<html-stem>.pptx)",
    )
    args = ap.parse_args()

    if not args.html.exists():
        sys.exit(f"missing input: {args.html}")

    if args.output is None:
        out_pptx = Path("build") / f"{args.html.stem}.pptx"
    else:
        out_pptx = args.output
    png_dir = out_pptx.parent / "png" / args.html.stem

    print(f"[1/2] Rendering {args.html.name} → PNG …")
    pngs = render_slides(args.html, png_dir)
    print(f"[2/2] Building PPTX with {len(pngs)} slides → {out_pptx} …")
    build_pptx(pngs, out_pptx)
    print("Done.")


if __name__ == "__main__":
    main()
